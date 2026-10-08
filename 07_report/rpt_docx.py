# -*- coding: utf-8 -*-
"""
07_report / rpt_docx.py
-----------------------
Draws the report as a Word document with python-docx: navy ink throughout,
muted table rules, equations as images at their true size, figures with
numbered captions, page numbers in the footer.

Author: Akosa Samuel Onyejekwe (independent)

Units: lengths in centimetres or inches as named; type sizes in points.
"""
import io
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

import project_meta as meta
from rpt_blocks import column_widths, runs, strip_tags
from rpt_data import ACCENT, BAND, NAVY, RULE, SOFT

FACE = "Calibri"
PAGE_W_CM, PAGE_H_CM, MARGIN_CM = 21.0, 29.7, 2.2
TEXT_W_IN = (PAGE_W_CM - 2*MARGIN_CM)/2.54
BODY_PT = 10.5


def rgb(code: str) -> RGBColor:
    """A colour from its hex code."""
    return RGBColor.from_string(code.lstrip("#"))


def set_face(style_or_run, size: float, colour: str = NAVY, bold: bool | None = None) -> None:
    """Type face, size and colour of a style or a run."""
    font = style_or_run.font
    font.name, font.size, font.color.rgb = FACE, Pt(size), rgb(colour)
    if bold is not None:
        font.bold = bold
    rpr = style_or_run.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), FACE)
    for theme in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if fonts.get(qn(theme)) is not None:
            del fonts.attrib[qn(theme)]


def page_field(paragraph) -> None:
    """Put the page-number field in a paragraph."""
    run = paragraph.add_run()
    set_face(run, 8.5, SOFT)
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def new_document() -> Document:
    """An A4 document with the styles, properties, header and footer set."""
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Cm(PAGE_W_CM), Cm(PAGE_H_CM)
    sec.left_margin = sec.right_margin = Cm(MARGIN_CM)
    sec.top_margin = sec.bottom_margin = Cm(2.3)
    sec.different_first_page_header_footer = True
    set_face(doc.styles["Normal"], BODY_PT)
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.12
    for name, size in (("Heading 1", 16), ("Heading 2", 12.5), ("Title", 26)):
        set_face(doc.styles[name], size, bold=True)
        doc.styles[name].paragraph_format.keep_with_next = True
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(18)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(8)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(12)
    doc.styles["Heading 2"].paragraph_format.space_after = Pt(5)
    set_face(doc.styles["List Bullet"], BODY_PT)
    props = doc.core_properties
    props.author = props.last_modified_by = meta.AUTHOR
    props.title, props.subject = meta.TITLE, meta.SUBTITLE
    props.comments, props.keywords, props.category = "", meta.METHOD, "Case study"
    props.created = props.modified = datetime.fromisoformat(meta.STUDY_DATE_ISO)
    head = sec.header.paragraphs[0]
    set_face(head.add_run(meta.TITLE), 8.5, SOFT)
    foot = sec.footer.paragraphs[0]
    set_face(foot.add_run(meta.AUTHOR_FULL + "    page "), 8.5, SOFT)
    page_field(foot)
    foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    return doc


def write_runs(paragraph, text: str, size: float = BODY_PT, colour: str = NAVY, bold: bool = False) -> None:
    """Add marked-up text to a paragraph."""
    for chunk, tags in runs(text):
        run = paragraph.add_run(chunk)
        set_face(run, size, colour, bold=bold or "b" in tags)
        run.italic = "i" in tags
        run.font.subscript = "sub" in tags
        run.font.superscript = "sup" in tags


def borders(table, colour: str, edges: dict) -> None:
    """Set the rules of a table; `edges` maps top, bottom, insideH ... to a
    width in eighths of a point (0 for none)."""
    pr = table._tbl.tblPr
    old = pr.find(qn("w:tblBorders"))
    if old is not None:
        pr.remove(old)
    box = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        width = edges.get(edge, 0)
        el.set(qn("w:val"), "single" if width else "nil")
        el.set(qn("w:sz"), str(width))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), colour.lstrip("#"))
        box.append(el)
    pr.append(box)


def shade(cell, colour: str) -> None:
    """Fill a cell."""
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), colour.lstrip("#"))
    cell._tc.get_or_add_tcPr().append(shd)


def fixed_widths(table, widths_in: list) -> None:
    """Fix the column widths of a table."""
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in table.rows:
        for cell, w in zip(row.cells, widths_in, strict=True):
            cell.width = Inches(w)


def repeat_header(row) -> None:
    """Mark a table row to repeat at the top of every page."""
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(el)


def add_title(doc, block: dict) -> None:
    """The cover."""
    for _ in range(7):
        doc.add_paragraph()
    write_runs(doc.add_paragraph(), block["kind_line"], 11, ACCENT, bold=True)
    write_runs(doc.add_paragraph(), block["title"], 26, NAVY, bold=True)
    write_runs(doc.add_paragraph(), block["subtitle"], 14, SOFT)
    doc.add_paragraph()
    for line in (block["solver"], block["method"]):
        write_runs(doc.add_paragraph(), line, 11.5)
    for _ in range(4):
        doc.add_paragraph()
    write_runs(doc.add_paragraph(), block["author"], 12)
    write_runs(doc.add_paragraph(), block["date"], 12)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_note(doc, block: dict) -> None:
    """A statement in a tinted band."""
    table = doc.add_table(rows=1, cols=1)
    borders(table, ACCENT, dict(left=24))
    fixed_widths(table, [TEXT_W_IN])
    cell = table.rows[0].cells[0]
    shade(cell, BAND)
    write_runs(cell.paragraphs[0], block["text"])
    cell.paragraphs[0].paragraph_format.space_before = Pt(4)
    cell.paragraphs[0].paragraph_format.space_after = Pt(5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_equation(doc, block: dict) -> None:
    """An equation centred, with its number at the right."""
    table = doc.add_table(rows=1, cols=3)
    borders(table, RULE, {})
    side = 0.5
    fixed_widths(table, [side, TEXT_W_IN - 2*side, side])
    left, mid, right = table.rows[0].cells
    for cell in (left, mid, right):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.paragraphs[0].paragraph_format.space_before = Pt(3)
        cell.paragraphs[0].paragraph_format.space_after = Pt(6)
    mid.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    mid.paragraphs[0].add_run().add_picture(str(block["path"]), width=Inches(block["width_in"]))
    right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    write_runs(right.paragraphs[0], f"({block['number']})")


def add_figure(doc, block: dict, pics) -> None:
    """One image, or a row of images, with its caption."""
    paths, n = block["paths"], len(block["paths"])
    each = (TEXT_W_IN*block["width"] - 0.25*(n - 1))/n
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.keep_with_next = True
    par.paragraph_format.space_before = Pt(6)
    par.paragraph_format.space_after = Pt(3)
    for i, path in enumerate(paths):
        w, h = pics.fit(path, each, block["max_h_in"])
        if i:
            par.add_run("  ")
        par.add_run().add_picture(io.BytesIO(pics.prepared(path, w, "docx")[1]), width=Inches(w), height=Inches(h))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(12)
    write_runs(cap, block["caption"], 9, SOFT)


def add_table(doc, block: dict) -> None:
    """A table with its caption above it."""
    cap = doc.add_paragraph()
    cap.paragraph_format.keep_with_next = True
    cap.paragraph_format.space_before = Pt(6)
    cap.paragraph_format.space_after = Pt(3)
    write_runs(cap, block["caption"], 9, SOFT)
    header, rows, size = block["header"], block["rows"], block["size"]
    widths = None
    while widths is None:
        per_char = size*0.0075                       # inches per character, a generous mean
        plain = [[strip_tags(c) for c in row] for row in rows] if block.get("rich") else rows
        widths = column_widths(header, plain, TEXT_W_IN, lambda text, c=per_char: len(text)*c)
        if widths is None and size <= 6.0:
            widths = [TEXT_W_IN/len(header)]*len(header)
        elif widths is None:
            size -= 0.5
    scale = TEXT_W_IN/sum(widths)
    table = doc.add_table(rows=len(rows) + 1, cols=len(header))
    borders(table, RULE, dict(top=8, bottom=8, insideH=4))
    fixed_widths(table, [w*scale for w in widths])
    repeat_header(table.rows[0])
    for j, text in enumerate(header):
        fill_cell(table.rows[0].cells[j], text, size, bold=True)
        shade(table.rows[0].cells[j], BAND)
    for i, row in enumerate(rows, start=1):
        for j, text in enumerate(row):
            fill_cell(table.rows[i].cells[j], text, size, rich=block.get("rich", False))
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def fill_cell(cell, text: str, size: float, bold: bool = False, rich: bool = False) -> None:
    """Text in a table cell: plain, or with the inline markup when `rich`."""
    par = cell.paragraphs[0]
    par.paragraph_format.space_before = Pt(1.5)
    par.paragraph_format.space_after = Pt(1.5)
    par.paragraph_format.line_spacing = 1.0
    if rich:
        write_runs(par, text, size, bold=bold)
    else:
        set_face(par.add_run(text), size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_text(doc, block: dict) -> None:
    """A heading, paragraph or bullet."""
    kind = block["kind"]
    if kind in ("h1", "h2"):
        par = doc.add_heading(level=int(kind[1]))
        write_runs(par, block["text"], 16 if kind == "h1" else 12.5, bold=True)
    elif kind == "bullet":
        write_runs(doc.add_paragraph(style="List Bullet"), block["text"])
    else:
        write_runs(doc.add_paragraph(), block["text"])


def render_report(report, pics, path: Path) -> dict:
    """Write the report as a Word document; returns counts of what it holds."""
    doc = new_document()
    for block in report.blocks:
        kind = block["kind"]
        if kind == "title":
            add_title(doc, block)
        elif kind == "note":
            add_note(doc, block)
        elif kind == "eq":
            add_equation(doc, block)
        elif kind == "fig":
            add_figure(doc, block, pics)
        elif kind == "table":
            add_table(doc, block)
        elif kind == "break":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        else:
            add_text(doc, block)
    doc.save(str(path))
    words = sum(len(strip_tags(b.get("text", "")).split()) for b in report.blocks)
    return dict(blocks=len(report.blocks), words_in_text_blocks=words)
