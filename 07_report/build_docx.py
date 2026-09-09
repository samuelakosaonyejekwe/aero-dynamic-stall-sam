# -*- coding: utf-8 -*-
"""
07_report / build_docx.py
-------------------------
Assembles the COMPLETE case study into case.docx — background, problem
statement, full governing equations, novelty/contribution, competitive
positioning, ALL input data, ALL generated output (CSV tables, metrics, every
graph/curve/contour/chart and engineering drawing), and all validation /
calibration sources.

Author: Akosa Samuel Onyejekwe (independent).  No third-party attribution.
"""
import glob
from pathlib import Path
import json
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
import sys; sys.path.insert(0, str(ROOT))
from project_meta import AUTHOR, STUDY_DATE, METHOD  # single source of truth
EQDIR = ROOT/"07_report"/"_equations"; EQDIR.mkdir(exist_ok=True)
plt.rcParams["mathtext.fontset"] = "cm"   # Computer-Modern: standard math look
_eqi = [0]
INK = RGBColor(0x1f, 0x33, 0x50)
ACC = RGBColor(0x2e, 0x6f, 0xb7)

doc = Document()
# ---- base styles (no black: body text deep navy) ----
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
st.font.color.rgb = INK
for h, sz in [("Heading 1", 16), ("Heading 2", 13.5), ("Heading 3", 12)]:
    s = doc.styles[h]; s.font.color.rgb = ACC; s.font.size = Pt(sz)

# ---------------------------------------------------------------- helpers
def P(text="", size=11, bold=False, italic=False, align=None, color=INK, space=6):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = bold; r.italic = italic
    r.font.size = Pt(size); r.font.color.rgb = color
    if align: p.alignment = align
    p.paragraph_format.space_after = Pt(space)
    return p

def H(text, lvl=1):
    doc.add_heading(text, level=lvl)

def EQ(latex, scale=1.0):
    """Render a LaTeX equation to an image (standard math typesetting, normal
    size) and insert it centred in the document."""
    _eqi[0] += 1
    fn = EQDIR/f"eq_{_eqi[0]:03d}.png"
    fig = plt.figure(figsize=(0.01, 0.01))
    fig.text(0, 0, f"${latex}$", fontsize=12, color=(0.122, 0.20, 0.314))
    fig.savefig(fn, dpi=200, transparent=True, bbox_inches="tight", pad_inches=0.03)
    plt.close(fig)
    from PIL import Image
    w, _h = Image.open(fn).size
    # rendered at 12 pt; embed at TRUE size (w/200 in) so on-page font == 12 pt
    # (matches the 11 pt body text); only very long lines are capped to column.
    win = min(6.0, (w/200.0)*scale)
    doc.add_picture(str(fn), width=Inches(win))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[-1].paragraph_format.space_after = Pt(4)
    doc.paragraphs[-1].paragraph_format.space_before = Pt(4)

def bullet(text):
    p = doc.add_paragraph(style="List Bullet"); r = p.add_run(text)
    r.font.size = Pt(11); r.font.color.rgb = INK

def caption(text):
    p = doc.add_paragraph()
    r = p.add_run(text); r.italic = True; r.font.size = Pt(9.5); r.font.color.rgb = ACC
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)

def add_image(path, width=6.3, cap=None):
    path = Path(path)
    if not path.exists(): return
    doc.add_picture(str(path), width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if cap: caption(cap)

# Cells with no value are rendered as an em dash. Writing str(NaN) put the bare
# string "nan" into 14 table cells of the consolidated report -- five of them the
# rotor parameters of Case A, which has no rotor because it is a wind-tunnel rig,
# and one the cycle-1 convergence residual, which has no predecessor to difference
# against. "nan" reads as a computation that failed; neither of those did.
NA_DASH = "\u2014"

def _fmt_cell(v):
    try:
        if v is None or (isinstance(v, float) and v != v):
            return NA_DASH
    except Exception:
        pass
    t = str(v)
    return NA_DASH if t.strip().lower() in ("nan", "none", "") else t


def add_table_from_df(df, max_rows=60, max_cols=12, note=None):
    d = df.copy()
    if d.shape[1] > max_cols:
        d = d.iloc[:, :max_cols]
    truncated = len(d) > max_rows
    if truncated: d = d.head(max_rows)
    has_na = any(_fmt_cell(v) == NA_DASH for row in d.itertuples(index=False) for v in row)
    t = doc.add_table(rows=1, cols=len(d.columns))
    t.style = "Light Grid Accent 1"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, c in enumerate(d.columns):
        cell = t.rows[0].cells[j]; cell.text = str(c)
        for r in cell.paragraphs[0].runs: r.bold = True; r.font.size = Pt(9)
    for _, row in d.iterrows():
        cells = t.add_row().cells
        for j, c in enumerate(d.columns):
            cells[j].text = _fmt_cell(row[c])
            for r in cells[j].paragraphs[0].runs: r.font.size = Pt(8.5)
    msg = []
    if truncated: msg.append(f"showing first {max_rows} of {len(df)} rows")
    if df.shape[1] > max_cols: msg.append(f"first {max_cols} of {df.shape[1]} columns")
    if has_na: msg.append(f"{NA_DASH} = not applicable to that case / no value defined")
    if note: msg.append(note)
    if msg: caption("Table: " + "; ".join(msg))
    else: doc.add_paragraph().paragraph_format.space_after = Pt(8)

def add_csv(path, **kw):
    try:
        add_table_from_df(pd.read_csv(path), **kw)
    except Exception as e:
        P(f"[could not render {Path(path).name}: {e}]", italic=True)

# ================================================================ TITLE
P("INDUSTRIAL CASE STUDY", size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, color=ACC, space=2)
P("Prediction of Dynamic Stall on a Helicopter Main-Rotor Retreating Blade",
  size=20, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space=2)
P("using the UNISTALL™ Universal Unsteady-Aerodynamics & Dynamic-Stall Solver",
  size=13, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space=14)
P(f"Core method: {METHOD}",
  size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space=2)
P(f"Author: {AUTHOR} (independent)", size=12, bold=True,
  align=WD_ALIGN_PARAGRAPH.CENTER, space=2)
P(f"Date: {STUDY_DATE}", size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space=2)
P("Validated against published NACA 0012 static & dynamic-stall data",
  size=10.5, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space=2)
doc.add_page_break()

# ================================================================ EXEC SUMMARY
H("Executive Summary", 1)
P("This case study demonstrates the prediction of dynamic stall — the single most "
  "important unsteady-aerodynamic phenomenon limiting the flight envelope of "
  "edgewise rotors — on the retreating blade of a medium utility helicopter, using "
  "the UNISTALL™ universal solver. The solver couples a unified state-space "
  "indicial–Beddoes (UIBS) reduced-order model with an integrated potential-flow "
  "field-reconstruction module and a compressible thermal module, producing the full "
  "engineering output set (unsteady load loops, surface pressures, pressure / "
  "velocity / vorticity / temperature fields, vectors, 3-D response surfaces and "
  "metrics) at a tiny fraction of the cost of time-accurate CFD.")
# headline numbers are READ from the solution, never restated, so this paragraph
# can never drift away from what the solver actually produced
_mA = pd.read_csv(ROOT/"05_solution"/"metrics_A_validation.csv").set_index("metric")["value"]
_vs = pd.read_csv(ROOT/"06_postprocessing"/"validation"/"validation_realdata_summary.csv"
                  ).set_index("metric")["value"]
P("The solver is calibrated per case to a published NACA 0012 static polar and "
  "validated against the classical NACA 0012 dynamic-stall experiments of "
  "McAlister, Carr & McCroskey. Static errors are below 1 %. The dynamic constants "
  f"are calibrated on ONE measured loop and then frozen to predict {_vs['NACA0012 held-out frames']} "
  f"held-out loops, over which the mean peak-lift error is {_vs['mean |CLmax| error [%]']} % and "
  f"the mean moment-break error {_vs['mean |CMmin| error [abs]']} (§12.2). Headline "
  "results for the matched validation point (M = 0.30, k = 0.10, α = 10° ± 10°): "
  f"dynamic C_L,max = {_mA['CL_max_dynamic']} "
  f"({100*(float(_mA['dynamic_overshoot_ratio'])-1):.0f} % overshoot above static), "
  f"C_M,c/4 break to {_mA['CM_min(c/4)']}, C_D,max = {_mA['CD_max']}.")
doc.add_page_break()

# ================================================================ 1 BACKGROUND
H("1. Background", 1)
P("Dynamic stall occurs when an aerofoil is driven rapidly beyond its static stall "
  "angle. The boundary layer separation is delayed by the unsteady motion, a strong "
  "leading-edge (dynamic-stall) vortex forms and is shed, and the section briefly "
  "develops lift far in excess of its static maximum. As the vortex convects over "
  "the chord and past the trailing edge it produces a large nose-down pitching-moment "
  "excursion (the 'moment break') followed by a sudden loss of lift and a hysteretic "
  "reattachment. The resulting cyclic loads drive rotor vibration, control loads, "
  "fatigue and stall flutter.")
P("On a helicopter in fast forward flight the retreating blade operates at low "
  "dynamic pressure and high pitch to balance the advancing side; near the azimuth "
  "ψ = 270° it is forced through exactly this rapid pitch-up and stalls "
  "dynamically once per revolution. Retreating-blade dynamic stall therefore sets "
  "the maximum forward speed and maximum rotor thrust of every conventional "
  "helicopter — making its accurate, cheap prediction a central industrial problem.")
P("High-fidelity CFD (URANS / DES / LES) can resolve dynamic stall but at a cost of "
  "hours-to-days per operating point and with strong sensitivity to turbulence and "
  "transition modelling. Reduced-order semi-empirical models (Leishman–Beddoes, "
  "ONERA, Gormont, Snel, Øye) are the industrial workhorses for rotor and wind-"
  "turbine loads because they run in milliseconds. UNISTALL™ extends this lineage.")

# ================================================================ 2 PROBLEM
H("2. Problem Statement", 1)
P("Given a rotor blade section undergoing a prescribed unsteady pitch schedule at a "
  "specified Mach and reduced frequency, predict — quickly, robustly and with "
  "quantified accuracy — the complete unsteady aerodynamic response, including:")
for b in ["the unsteady lift, drag and pitching-moment hysteresis loops;",
          "the dynamic-stall onset angle, lift overshoot and moment-break magnitude;",
          "the trailing-edge separation history and the dynamic-stall vortex trajectory;",
          "the surface pressure distribution and the 2-D pressure / velocity / vorticity field;",
          "the compressible static and recovery (skin) temperature field;",
          "the aerodynamic damping (stall-flutter indicator); and",
          "the sensitivity of the above to mean incidence and reduced frequency."]:
    bullet(b)
P("The prediction must be validated against credible published experimental data and "
  "calibrated to the specific section, and must be orders of magnitude cheaper than "
  "time-accurate CFD so that it can be embedded in design sweeps and loads loops.")

# ================================================================ 3 CASE DEF
H("3. Industrial Case-Study Definition", 1)
# every number in this section is READ from 03_model_setup, so the prose cannot
# drift away from the conditions that were actually solved
_fl = pd.read_csv(ROOT/"03_model_setup"/"flow_conditions.csv").set_index("parameter")
_kn = pd.read_csv(ROOT/"03_model_setup"/"kinematics.csv").set_index("case_id")
_A, _B = _fl["case_A_validation"], _fl["case_B_application"]
_kA, _kB = _kn.loc["A_validation_rig"], _kn.loc["B_application_rotor"]
P("Reference aircraft: a generic medium utility helicopter (\u201cCS-MUH reference\u201d), "
  f"4-bladed main rotor of radius R = {float(_B['rotor_radius_R']):.2f} m, blade chord "
  f"{float(_B['chord_c']):.3f} m, NACA 0012 section, tip speed \u03a9R \u2248 "
  f"{float(_B['rotor_tip_speed_OmegaR']):.0f} m/s. Forward flight at advance ratio "
  f"\u03bc = {float(_B['advance_ratio_mu']):.2f}. The analysis section is the retreating "
  f"blade at r/R = {float(_B['radial_station_r_R']):.2f}, where the section velocity is "
  f"\u03a9R(r/R \u2212 \u03bc) = {float(_B['freestream_velocity_U']):.1f} m/s. A "
  "geometrically-matched NACA 0012 wind-tunnel oscillating-aerofoil rig is used as the "
  "validation configuration.")
P("The free-stream state is over-determined \u2014 (p, T, \u03c1), (M, U, a) and, for "
  "Case B, the rotor kinematics all describe the same flow \u2014 so only the "
  "independent quantities are specified and everything else is derived: "
  "a = \u221a(\u03b3RT), \u03c1 = p/(RT), Re_c = \u03c1Uc/\u00b5, with U = M\u00b7a "
  "for the rig and U = \u03a9R(r/R \u2212 \u03bc) for the blade station.",
  italic=True, size=10)
P("Two configurations are solved:", bold=True)
bullet(f"Case A \u2014 validation rig: NACA 0012, M = {float(_A['freestream_mach_M']):.2f}, "
       f"k = {float(_kA['reduced_freq_k']):.2f}, "
       f"\u03b1 = {float(_kA['alpha_mean_deg']):.0f}\u00b0 \u00b1 "
       f"{float(_kA['alpha_amp_deg']):.0f}\u00b0 "
       "(matches the McAlister/McCroskey deep dynamic-stall test point).")
bullet(f"Case B \u2014 application: retreating-blade section "
       f"r/R = {float(_B['radial_station_r_R']):.2f}, "
       f"M = {float(_B['freestream_mach_M']):.2f}, k = {float(_kB['reduced_freq_k']):.3f}, "
       f"\u03b1 = {float(_kB['alpha_mean_deg']):.0f}\u00b0 \u00b1 "
       f"{float(_kB['alpha_amp_deg']):.0f}\u00b0 (1/rev feathering).")
P("Full operating conditions (solver input):", bold=True)
add_csv(ROOT/"03_model_setup"/"flow_conditions.csv")
P("Blade pitch kinematics (solver input):", bold=True)
add_csv(ROOT/"03_model_setup"/"kinematics.csv")

# ================================================================ 4 THEORY/EQNS
H("4. Solver Theory and Governing Equations", 1)
P("UNISTALL™ marches the UIBS state-space model in semichord (reduced) time "
  "s = 2Ut/c. The model superposes attached-flow indicial loads, a Kirchhoff "
  "trailing-edge separation model with two time lags, and a leading-edge dynamic-"
  "stall vortex. All equations actually implemented in the solver are listed below.")

H("4.1 Reduced time and compressibility", 2)
EQ(r"s = \frac{2U}{c}\,t,\qquad \Delta s = \frac{2U\,\Delta t}{c},\qquad \beta=\sqrt{1-M^{2}}")
H("4.2 Attached-flow circulatory loads (Beddoes two-lag indicial)", 2)
P("Three-quarter-chord angle of attack (pitch about the quarter chord):")
EQ(r"\alpha_{3/4}=\alpha+\frac{c}{2U}\,\dot\alpha")
EQ(r"X_{1}^{\,n}=X_{1}^{\,n-1}e^{-b_{1}\beta^{2}\Delta s}+A_{1}\,\Delta\alpha_{3/4}\,e^{-b_{1}\beta^{2}\Delta s/2}")
EQ(r"X_{2}^{\,n}=X_{2}^{\,n-1}e^{-b_{2}\beta^{2}\Delta s}+A_{2}\,\Delta\alpha_{3/4}\,e^{-b_{2}\beta^{2}\Delta s/2}")
EQ(r"\alpha_{E}=\alpha_{3/4}-X_{1}-X_{2},\qquad C_{N}^{C}=C_{N\alpha}\,\alpha_{E}")
# Values taken FROM solver_config.json, not restated. They were written here as
# literals, a second copy of numbers the solver reads from the config, and a
# change to one would have left the report quoting the other.
_ic = json.load(open(ROOT/"03_model_setup"/"solver_config.json"))["indicial_circulatory"]
EQ(r"(A_{1},A_{2},b_{1},b_{2})=(%.2f,\;%.2f,\;%.2f,\;%.2f)"
   % (_ic["A1"], _ic["A2"], _ic["b1"], _ic["b2"]))
H("4.3 Non-circulatory (impulsive / added-mass) loads", 2)
EQ(r"K_{\alpha}=\frac{0.75}{1-M+\pi\beta M^{2}(A_{1}b_{1}+A_{2}b_{2})},\qquad "
   r"T_{I}=\frac{K_{\alpha}c}{U}=\frac{K_{\alpha}c}{M\,a}")
P("The impulsive time constant as implemented, written out because it is a "
  "stated departure rather than a silent one. In semichord time Δt/T_I = "
  "Δs/(2K_α), so the impulsive lag is a fixed 2K_α semichords and the march "
  "stays chord- and speed-independent. This is NOT the classical "
  "Leishman–Beddoes T_I = c/a: it is larger by 1/M (3.3× at M = 0.30), and the "
  "amplitude 4K_αc/(UM) below carries the same extra 1/M against the classical "
  "4K_αc/U. It is the convention the calibrated constants of §10 were fitted "
  "with; swapping both terms to the classical scaling and re-running the five "
  "real NACA 0012 frames of §12.2 with those same constants moves the held-out "
  "mean peak-lift error from 1.7 % to 3.0 %, so the implemented form is kept "
  "and documented rather than silently reinterpreted. This equation previously "
  "read T_I = K_αc/a, which the solver has never computed.",
  italic=True, size=10)
EQ(r"D_{I}^{\,n}=D_{I}^{\,n-1}e^{-\Delta t/T_{I}}+\left(\dot\alpha_{3/4}^{\,n}-\dot\alpha_{3/4}^{\,n-1}\right)e^{-\Delta t/2T_{I}}")
EQ(r"C_{N}^{I}=\frac{4K_{\alpha}c}{UM}\left(\dot\alpha_{3/4}-D_{I}\right),\qquad C_{N}^{P}=C_{N}^{C}+C_{N}^{I}")
H("4.4 Trailing-edge separation (pressure & boundary-layer lags, Kirchhoff)", 2)
EQ(r"D_{p}^{\,n}=D_{p}^{\,n-1}e^{-\Delta s/T_{p}}+\left(C_{N}^{P,n}-C_{N}^{P,n-1}\right)e^{-\Delta s/2T_{p}}")
EQ(r"C_{N}'=C_{N}^{P}-D_{p},\qquad \alpha_{f}=\frac{C_{N}'}{C_{N\alpha}},\qquad f'=f_{\mathrm{static}}(\alpha_{f})")
EQ(r"D_{f}^{\,n}=D_{f}^{\,n-1}e^{-\Delta s/T_{f}}+\left(f'_{n}-f'_{n-1}\right)e^{-\Delta s/2T_{f}}")
P("Dynamic separation point and the Kirchhoff/Helmholtz normal force:")
EQ(r"f''=f'-D_{f},\qquad C_{N}^{f}=C_{N\alpha}\,\alpha_{E}\left(\frac{1+\sqrt{f''}}{2}\right)^{2}+C_{N}^{I}")
P("Calibration of the static separation point by inversion of the Kirchhoff relation:")
EQ(r"f_{\mathrm{static}}=\left(2\sqrt{\frac{C_{N}^{\,st}}{C_{N\alpha}\,\alpha}}-1\right)^{2}")
H("4.5 Leading-edge dynamic-stall vortex", 2)
P("Vortex shedding is triggered when the delayed normal force reaches the critical "
  "value on the up-stroke; the vortex is then fed and convected:")
EQ(r"C_{N}'\geq C_{N1}\ \ (\dot\alpha>0),\qquad C_{v}=B_{v}\,C_{N}^{C}\left[\,1-\left(\frac{1+\sqrt{f''}}{2}\right)^{2}\right]")
EQ(r"\tau_{v}\leq T_{vl}:\quad C_{N}^{v,n}=C_{N}^{v,n-1}e^{-\Delta s/T_{v}}+\left(C_{v,n}-C_{v,n-1}\right)e^{-\Delta s/2T_{v}}")
EQ(r"\tau_{v}>T_{vl}:\quad C_{N}^{v,n}=C_{N}^{v,n-1}e^{-\Delta s/T_{v}}")
H("4.6 Total section loads", 2)
EQ(r"C_{N}=C_{N}^{f}+C_{N}^{v},\qquad C_{C}=\eta\,C_{N\alpha}\,\alpha_{E}^{2}\sqrt{f''}")
EQ(r"C_{L}=C_{N}\cos\alpha+C_{C}\sin\alpha,\qquad C_{D}=C_{N}\sin\alpha-C_{C}\cos\alpha+C_{D0}")
EQ(r"C_{M}=C_{M0}+\left[k_{0}+k_{1}(1-f'')+k_{2}\sin(\pi f''^{\,\kappa})\right]C_{N}^{f}-\mathrm{CP}_{v}\,C_{N}^{v}")
EQ(r"\mathrm{CP}_{v}=a_{v}\left(1-\cos\!\left(\pi\,\mathrm{min}"
  r"\!\left(\tau_{v}/T_{vl},\,1\right)\right)\right)")
P("Aerodynamic damping (cyclic work), and its normalisation by the area of the "
  "bounding box of the C_M–α loop. Only the normalised form is a usable "
  "discriminator: for the near-cancelling figure-of-eight loops of light dynamic "
  "stall the raw Ξ is a small residual between two lobes of opposite sign.")
EQ(r"\Xi=-\oint C_{M}\,d\alpha,\qquad \hat\Xi=\frac{\Xi}"
  r"{\left(\Delta C_{M}\right)\left(\Delta\alpha\right)}")
P("A verdict is only reported outside a neutral band. That band is set by the model’s "
  "own demonstrated accuracy in this quantity — the mean discrepancy between the "
  "modelled and measured Ξ̂ over the real NACA 0012 loops of §12.2 — "
  "not by the time-step error, which is some 180 times smaller than the band "
  "(refining 720 → 5760 steps/cycle moves Ξ̂ by 0.0004, against a band of "
  "0.08). The measured "
  "spread and the band in force are both tabulated in §12.2.", italic=True, size=10)
H("4.7 Field reconstruction (pressure / velocity / vorticity)", 2)
EQ(r"\sum_{j}\sigma_{j}\left(\frac{\partial\phi_{j}}{\partial n}\right)_{i}=-\,\mathbf{U}_{\infty}\!\cdot\mathbf{n}_{i}")
P("Bound circulation from Kutta–Joukowski, carried as a uniform vortex sheet on "
  "the body surface and solved together with the source panels, so the trailing edge "
  "satisfies the Kutta condition to within the residual quoted in §4.9:")
EQ(r"\Gamma=\frac{1}{2}C_{L}U c,\qquad \gamma=\frac{\Gamma}{\oint \mathrm{d}s}")
EQ(r"V_{\theta}=\frac{\Gamma_{v}}{2\pi r}\left(1-e^{-r^{2}/r_{c}^{2}}\right),\qquad \Gamma_{v}\propto C_{N}^{v}")
P("Surface and field pressure coefficient. There is deliberately NO "
  "Prandtl–Glauert factor on this expression, and the equation used to carry "
  "one: compressibility has already entered through β in the indicial march "
  "that produced the C_L, and Γ = ½ C_L U c hands that C_L to the "
  "reconstruction, so a second 1/√(1−M²) here multiplies the reconstructed load "
  "again — 4.8 % at M = 0.3, one of the three errors dissected in §4.9. The "
  "Bernoulli term holds only where the flow is irrotational, so inside the "
  "dynamic-stall vortex core it is corrected to radial equilibrium "
  "(dp/dr = ρvθ²/r); the correction decays to zero outside the core, "
  "where radial equilibrium and Bernoulli agree:")
EQ(r"C_{p}=1-\left(\frac{V}{U}\right)^{2}"
  r"+\Delta C_{p}^{\,\mathrm{core}}(r)")
EQ(r"\Delta C_{p}^{\,\mathrm{core}}(r)=-\frac{2}{U^{2}}\int_{r}^{\infty}"
  r"\frac{V_{\theta}^{2}}{r'}\,\mathrm{d}r' + \left(\frac{V_{\theta}}{U}\right)^{2}")
H("4.8 Compressible thermal module", 2)
EQ(r"T_{0}=T_{\infty}\left(1+\frac{\gamma-1}{2}M^{2}\right)")
EQ(r"T=T_{0}-\frac{V^{2}}{2c_{p}},\qquad T_{r}=T_{0}-(1-r)\frac{V^{2}}{2c_{p}},\qquad r=Pr^{1/3}")
EQ(r"M_{\mathrm{local}}=\frac{V}{\sqrt{\gamma R T}}")
P("Air thermodynamic properties (solver input). c_p and the recovery factor are "
  "derived from γ, R and Pr rather than quoted independently:", bold=True)
add_csv(ROOT/"03_model_setup"/"material_thermo_properties.csv")

H("4.9 Accuracy of the field reconstruction", 2)
P("The reconstruction is a potential field and is reported as qualitative. Its "
  "checkable properties are measured on every run rather than asserted:")
_mA_c = pd.read_csv(ROOT/"05_solution"/"metrics_A_validation.csv").set_index("metric")["value"]
_mB_c = pd.read_csv(ROOT/"05_solution"/"metrics_B_application.csv").set_index("metric")["value"]
bullet("Closure. Integrating the reconstructed surface C_p must return the C_L the "
       "reconstruction was given (Γ = ½ C_L U c). At peak incidence it returns it to "
       f"{_mA_c['Cp_closure_error_pct']} % (Case A) and {_mB_c['Cp_closure_error_pct']} % "
       "(Case B), published as Cp_closure_error_pct in metrics_*.csv. That is a "
       "single instant and is not a bound on the cycle, which an earlier revision "
       "implied it was. What is measured over the whole cycle is the worst ABSOLUTE "
       "residual, in C_L counts: "
       f"{_mA_c['Cp_closure_worst_dCL_cycle']} (Case A) and "
       f"{_mB_c['Cp_closure_worst_dCL_cycle']} (Case B), i.e. "
       f"{_mA_c['Cp_closure_worst_dCL_pct_of_CLmax']} % and "
       f"{_mB_c['Cp_closure_worst_dCL_pct_of_CLmax']} % of each case's own C_L,max "
       "(Cp_closure_worst_dCL_cycle and Cp_closure_worst_dCL_pct_of_CLmax). It is "
       "reported in C_L counts rather than as a worst instantaneous percentage "
       "because the cycle passes through C_L = 0.09, where a residual of 0.0014 "
       "reads as +1.6 % purely from the small denominator.")
bullet("Kutta condition. The trailing-edge C_p jump is NOT a residual that can be "
       "driven to zero, and is no longer presented as one. It is linear in the "
       "imposed C_L and passes through zero exactly at the inviscid attached "
       "circulation, published beside it as CL_kutta_inviscid "
       f"({_mA_c['CL_kutta_inviscid']} for Case A); imposing that value drives the "
       "jump to about 0.001, and the measured ratio jump/|C_L − C_L,Kutta| is "
       "1.91–2.02 across α = 2–19°, falling monotonically with incidence. The "
       "reconstruction is instead handed the indicial C_L, "
       "which during dynamic stall departs from the attached value deliberately, so a "
       "body carrying a non-Kutta circulation must show a jump. Over the cycle phases "
       f"written out it reaches {_mA_c['Cp_TE_jump_max_over_phases']} (Case A) and "
       f"{_mB_c['Cp_TE_jump_max_over_phases']} (Case B), published as "
       "Cp_TE_jump_max_over_phases: a measure of how far the modelled flow is from "
       "attached, not an error.")
bullet("Dynamic-stall-vortex core depth, reported rather than tuned away. The "
       "suction at the centre of the reconstructed vortex is published as "
       f"Cp_DSV_core_min: {_mA_c['Cp_DSV_core_min']} (Case A) and "
       f"{_mB_c['Cp_DSV_core_min']} (Case B), where a measured deep-stall core is "
       "usually nearer −3 to −6. Deepening it means shrinking the core radius and "
       "raising the circulation factor together, and neither constant can be "
       "calibrated from anything this study ships: the experimental frames carry only "
       "integrated C_L, C_D and C_M against incidence, with no surface-pressure or "
       "field data to fit a core size to. Both are therefore named constants in the "
       "solver (DSV_GAMMA_FACTOR, DSV_CORE_RADIUS_CHORDS) and the resulting depth is "
       "published as a number, so the shallowness is checkable rather than an "
       "adjective.")
P("This closure figure was −12.4 % in an earlier revision of this study, and the "
  "explanation recorded for it — the C_p clip at −8, together with the claim that the "
  "deficit did not vanish under refinement — was tested and found false on both "
  "counts: moving the clip to −10⁹ changed the closure by 0.00 points, and refining "
  "160 → 1280 panels moved it monotonically from −11.5 % to −5.2 %. The three real "
  "causes were all in the evaluation rather than the physics. The vortex sheet's own "
  "tangential contribution (−γ/2, the velocity jump across a sheet) was omitted, worth "
  "−48.7 % of the lift at the wall on its own. C_p was evaluated at the panel "
  "end-points, stepped 0.015c off the wall, rather than at the control points where "
  "flow tangency is actually imposed — an offset that was masking the missing term "
  "rather than avoiding it, since the error grew toward −48.7 % as the probe "
  "approached the surface. And a Prandtl–Glauert factor was applied to a C_p whose "
  "circulation already carried compressibility, inflating the load a further 4.8 % at "
  "M = 0.3. With all three corrected the closure converges: refining 160 → 1280 panels "
  "now drives it monotonically to −0.04 %, which is what Blasius requires of an exact "
  "potential solution and is the check that the formulation is right rather than "
  "merely closer. Nothing is clipped; the field nonetheless bottoms out near C_p = "
  "−5.2 because the near-wall ring carrying the leading-edge peak is masked, so the "
  "contour plots understate the surface suction (about −15) by roughly three times. "
  "Nothing in the reconstruction knows about separation; the reported loads come from "
  "the UIBS core and do not depend on any of it.",
  italic=True, size=10)

# ================================================================ 5 NOVELTY
H("5. Novelty and Contribution to Knowledge", 1)
for b in [
 "Unified state-space form (UIBS): the attached-flow, separation and vortex sub-"
 "models are recast as a single exponential-recurrence state machine in semichord "
 "time, giving a numerically robust, unconditionally-marching scheme with no "
 "iteration and no mesh.",
 "Automatic per-case calibration: the trailing-edge separation function f(α) is "
 "derived directly by inverting the Kirchhoff relation against a single measured "
 "static polar, so the model self-calibrates to any section without hand-tuning.",
 "Integrated reduced-order field reconstruction: the model's circulation and vortex "
 "states drive a source-panel + bound-sheet + Lamb–Oseen reconstruction that yields "
 "CFD-like pressure, velocity, vorticity and streamline fields at reduced-order cost "
 "— bridging the visualisation gap between ROMs and CFD.",
 "Embedded compressible thermal module: static and recovery (skin) temperature "
 "fields are produced alongside the loads — rarely available from dynamic-stall ROMs "
 "and directly useful for blade thermal/anti-icing and material assessment.",
 "Design-space throughput: a sub-second per-cycle march (measured: see §6) "
 "enables full response surfaces C_L,max(α_mean, k) and loads at interactive "
 "speed, which time-accurate CFD cannot deliver in a design loop."]:
    bullet(b)
P("These elements — the unified self-calibrating state-space core, the coupled field "
  "reconstruction, and the integrated thermal module in one solver — constitute the "
  "novel, patentable contribution of UNISTALL™.", italic=True)

# ================================================================ 6 COMPETITIVE
H("6. Competitive Positioning", 1)
P("The table contrasts UNISTALL™ with the main alternatives. Cost figures are order-"
  "of-magnitude; the quantified, substantiated claim is the validation accuracy in "
  "§12. High-fidelity CFD remains more general for novel geometries — UNISTALL™'s "
  "advantage is robustness, coupling and throughput at engineering accuracy.")
# The UNISTALL column asserted "< 1 s". It is now MEASURED rather than asserted,
# and the measurement confirms it: the march costs well under a second of CPU.
# The competing columns remain order-of-magnitude literature figures, which is
# what the paragraph above says they are.
_env = pd.read_csv(ROOT/"05_solution"/"runtime_environment.csv").set_index("property")["value"]
_cost = f"{float(_env['solve_cpu_time_s_case_A']):.2f} s CPU (measured)"
comp = pd.DataFrame([
 ["Cost per operating point (loads)", _cost, "hours–days", "days–weeks", "seconds"],
 ["Mesh required", "none (panel field opt.)", "yes (y+≈1)", "yes (fine)", "none"],
 ["Dynamic-stall loads", "yes (validated)", "yes", "yes", "yes"],
 ["2-D pressure/velocity field", "yes (reconstructed)", "yes", "yes", "no"],
 ["Temperature / recovery field", "yes", "with energy eqn", "with energy eqn", "no"],
 ["Auto-calibration from polar", "yes", "n/a", "n/a", "manual"],
 ["Response-surface sweeps", "yes (real-time)", "impractical", "impractical", "yes"],
 ["Robustness / no divergence", "high", "solver-dependent", "solver-dependent", "high"],
], columns=["Capability", "UNISTALL™ (UIBS)", "URANS CFD", "LES/DES",
            "Classic L-B / ONERA ROM"])
add_table_from_df(comp, max_rows=20, max_cols=6)
P(f"The UNISTALL™ figure is this study's own measured CPU time for the Case-A march "
  f"({int(float(_env['solve_cpu_ms_per_cycle_case_A']))} ms per cycle over "
  f"{_env['n_cycles']} cycles of {_env['steps_per_cycle']} steps, "
  f"{int(float(_env['solve_cpu_us_per_step']))} µs per step), taken on "
  f"Python {_env['python']} / {_env['processor']} and published in "
  "runtime_environment.csv. It is CPU time deliberately: wall time for identical "
  "code varied by a factor of seven on this machine with background load, so a "
  "wall-clock figure would say more about the machine than the solver. Writing "
  "the four reconstructed fields and four surface-C_p phases per case adds "
  "several seconds more. The other columns are order-of-magnitude figures from "
  "the literature, not measurements made here.", italic=True, size=10)
P("That figure is a property of this implementation, not of the method. Profiling "
  "the march puts about 60 % of it inside the static separation function, "
  "which is a vectorised SciPy/NumPy interpolation called once per step on a "
  "single scalar — per-call overhead, not arithmetic. A compiled or "
  "batch-evaluated implementation of the same equations would run far faster; "
  "the claim made here is only the CPU time this code actually took, which is "
  "the same measure as the table above and not a wall-clock figure.",
  italic=True, size=10)

# ================================================================ 7 GEOMETRY
H("7. Geometry", 1)
P(f"Primary section NACA 0012, model chord {float(_A['chord_c']):.2f} m (rig) / "
  f"{float(_B['chord_c']):.3f} m (blade). Surface coordinates are cosine-clustered at "
  "the leading and trailing edges.")
add_csv(ROOT/"01_geometry"/"section_geometry_summary.csv")
add_image(ROOT/"01_geometry"/"fig_geometry_profile.png", 6.3, "Fig. 7.1 NACA 0012 section profile and camber line.")
add_image(ROOT/"01_geometry"/"fig_geometry_thickness.png", 5.6, "Fig. 7.2 Thickness and camber distributions.")
P("Surface coordinate table (head):", bold=True)
add_csv(ROOT/"01_geometry"/"naca0012_coordinates.csv", max_rows=15)

# ================================================================ 8 MESH
H("8. Mesh", 1)
P("A body-fitted O-grid (the wall line wraps the whole surface; there is no wake "
  "cut) is generated for the panel field-reconstruction module and optional CFD "
  "hand-off, clustered to y+ ≈ 1 at the wall.")
add_csv(ROOT/"02_mesh"/"mesh_quality_metrics.csv")
add_image(ROOT/"02_mesh"/"fig_mesh_full.png", 5.4, "Fig. 8.1 Body-fitted O-grid (near field).")
add_image(ROOT/"02_mesh"/"fig_mesh_le_zoom.png", 4.6, "Fig. 8.2 Leading-edge boundary-layer clustering.")
add_image(ROOT/"02_mesh"/"fig_mesh_te_zoom.png", 4.6, "Fig. 8.3 Trailing-edge / near-wake clustering.")
add_image(ROOT/"02_mesh"/"fig_mesh_wall_spacing.png", 5.6, "Fig. 8.4 Wall-normal spacing law.")

# ================================================================ 9 SETUP/INPUT
H("9. Model Setup and Input Data", 1)
P("Calibration / validation static polar (published NACA 0012):", bold=True)
add_csv(ROOT/"03_model_setup"/"static_polar_reference.csv")

# ================================================================ 10 CALIB
H("10. Calibration", 1)
P("The static separation law is fitted to the published polar (inverse Kirchhoff); "
  "the dynamic constants are calibrated against a single measured oscillating-aerofoil "
  "loop (frame 9302) and then frozen. The per-frame errors that result are tabulated "
  "in §12.2 — no aggregate 'envelope' claim is made beyond those numbers.")
add_csv(ROOT/"06_postprocessing"/"validation"/"calibration_constants.csv")
add_image(ROOT/"06_postprocessing"/"plots"/"static_polar_calibration.png", 6.3,
          "Fig. 10.1 Calibrated static lift polar and separation function f(α).")

# ================================================================ 11 RESULTS
H("11. Solution — Engineering Outputs", 1)
for _i, (cs, lab) in enumerate([("A_validation", "Case A (validation rig)"),
                                ("B_application", "Case B (application rotor)")], start=1):
    H(f"11.{_i} {lab} — metrics", 2)
    add_csv(ROOT/"05_solution"/f"metrics_{cs}.csv")
P("Combined case summary:", bold=True)
add_csv(ROOT/"05_solution"/"summary_all_cases.csv")
_mA_d = pd.read_csv(ROOT/"05_solution"/"metrics_A_validation.csv").set_index("metric")["value"]
P("A note on the drag loop. The instantaneous C_D is negative over roughly a third of "
  f"the cycle (minimum {_mA_d['CD_min']} for Case A), which is visible in §11.3. That is not an "
  "error: on the downstroke the effective incidence lags one to two degrees above the "
  "geometric one, so the leading-edge suction term C_C cos α outweighs C_N sin α — the "
  "unsteady-thrust behaviour a chord-force model is built to reproduce. The invariant "
  "that must hold is the cycle mean, and it does: "
  f"{_mA_d['CD_cycle_mean']} for Case A, i.e. positive, so there is no net propulsion. "
  "CD_min and CD_cycle_mean are published alongside CD_max in metrics_*.csv.",
  italic=True, size=10)

H("11.3 Hysteresis loops", 2)
for cs in ["A_validation", "B_application"]:
    for v, nm in [("cl", "lift"), ("cd", "drag"), ("cm", "moment")]:
        add_image(ROOT/"06_postprocessing"/"plots"/f"hyst_{v}_{cs}.png", 4.7,
                  f"Hysteresis loop — {nm}, {cs.replace('_',' ')}.")
H("11.4 Time histories and internal states", 2)
for cs in ["A_validation", "B_application"]:
    add_image(ROOT/"06_postprocessing"/"plots"/f"timehist_loads_{cs}.png", 5.2,
              f"Unsteady load time histories, {cs.replace('_',' ')}.")
    add_image(ROOT/"06_postprocessing"/"plots"/f"states_{cs}.png", 5.8,
              f"UIBS internal states (separation & vortex), {cs.replace('_',' ')}.")
H("11.5 Surface pressure distributions", 2)
for cs in ["A_validation", "B_application"]:
    add_image(ROOT/"06_postprocessing"/"plots"/f"cp_distribution_{cs}.png", 5.6,
              f"Surface C_p at cycle phases, {cs.replace('_',' ')}.")
H("11.6 Convergence", 2)
add_image(ROOT/"06_postprocessing"/"plots"/"convergence_residuals.png", 5.4,
          "Cycle-to-cycle peak-C_L convergence.")
P("Convergence (Case A):", bold=True)
add_csv(ROOT/"05_solution"/"convergence"/"residuals_A_validation.csv")
P("Time-history sample (Case A, head):", bold=True)
add_csv(ROOT/"05_solution"/"time_history_A_validation.csv", max_rows=12)

# ================================================================ 12 VALIDATION
H("12. Validation against Published & Experimental Data", 1)
P("Validation has two layers: a static calibration check, and a dynamic "
  "validation against REAL digitised experimental loops using a strict "
  "calibrate-once / predict-the-rest protocol.")

H("12.1 Static validation (calibration check)", 2)
add_csv(ROOT/"06_postprocessing"/"validation"/"validation_static.csv")
add_image(ROOT/"06_postprocessing"/"validation"/"fig_validation_static_polar.png", 5.8,
          "Fig. 12.1 Static lift-polar validation (errors below 1 %).")

H("12.2 Dynamic validation against REAL digitised experimental loops", 2)
P("The dynamic constants are calibrated ONLY on one real NACA 0012 loop "
  "(frame 9302 = Case A, 10°±10°, M0.30, k0.10) and then FROZEN. The frozen "
  "model is used to PREDICT four held-out real NACA 0012 loops spanning "
  "light→deep stall and reduced frequency. Data: McAlister/Pucci/McCroskey/Carr "
  "(1982), NASA TM-84245, digitised via the open BL-DSM-JFS-2021 repository.")
P("Airfoil identity is CONFIRMED (not inferred) from that repository's "
  "load_frame.m mapping: frames 7019–14220 = NACA 0012, 24022–31310 = AMES-01, "
  "≥67000 = NLR-7301. All NACA 0012 frames below lie in 7019–14220; frame 25104 "
  "(AMES-01) is shown only as a labelled cross-check.", bold=True)
# Rendered as two explicit column subsets rather than "the first N columns".
# The CSV has 16 columns; the previous max_cols=13 silently dropped exactly the
# two the damping paragraph below points the reader at.
_nr = pd.read_csv(ROOT/"06_postprocessing"/"validation"/"validation_nasa_real.csv")
# TWO tables, conditions then errors, not one wide one. Adding peak_alpha_deg and
# within_static_calibration to the single 13-column table took it to 15 and
# squeezed the columns until the "airfoil" header broke mid-word -- which
# verify_invariants caught. Splitting on the natural seam (what was run vs how
# well it matched) gives 9 and 7 columns, both of which fit at full size.
add_table_from_df(_nr[["frame", "airfoil", "role", "M", "k", "alpha0_deg", "amp_deg",
                       "peak_alpha_deg", "within_static_calibration"]],
                  max_rows=8, max_cols=9,
                  note="conditions; airfoil identity source is the same for every row "
                       "(load_frame.m mapping, Pancini repo; data NASA TM-84245). "
                       "peak_alpha_deg = alpha0 + amplitude; within_static_calibration "
                       "is False where that peak exceeds the incidence the static "
                       "separation law is fitted to")
P("Errors against the same measured loops:", bold=True)
add_table_from_df(_nr[["frame", "RMS_CL", "RMS_CM", "CLmax_model", "CLmax_exp",
                       "CMmin_model", "CMmin_exp"]],
                  max_rows=8, max_cols=7)
P("Two of the four held-out frames extrapolate, and it costs accuracy. Frames 9217 "
  "and 9214 are 15° ± 10°, so they peak at 25° — five degrees past the "
  f"{float(_vs['static polar calibrated to [deg]']):.0f}° the static polar the separation law "
  "is fitted to is tabulated to. Every other part of this study flags that boundary: "
  "the response surface of §13 is bounded by it, the model polar carries a "
  "within_calibration column, the field-sampling incidences stop just inside it, and "
  "the calibration figures draw the extrapolated tail dashed. This validation did not, "
  "and the split is material — mean RMS C_L is "
  f"{_vs['mean RMS_CL, inside the calibration range']} over the two held-out frames "
  f"inside the calibration range against "
  f"{_vs['mean RMS_CL, peaking past it']} over the two beyond it, roughly double. The "
  "per-frame within_static_calibration column above, the split means in the summary "
  "table below, and the shaded band on Fig. 12.2 all record it rather than leaving a "
  "reader to notice that two of the loops run past the fit.", italic=True, size=10)
P("Normalised aerodynamic damping, model against the same measured loops "
  "(per-frame; the spread over these rows is what sets the neutral band):", bold=True)
add_table_from_df(_nr[["frame", "airfoil", "Xihat_model", "Xihat_exp"]],
                  max_rows=8, max_cols=4)
P("Held-out accuracy summary (frozen model vs real NACA 0012 experiment):", bold=True)
add_csv(ROOT/"06_postprocessing"/"validation"/"validation_realdata_summary.csv")
add_image(ROOT/"06_postprocessing"/"validation"/"fig_validation_nasa_real.png", 6.1,
          "Fig. 12.2 UNISTALL (calibrated on frame 9302 only) vs real digitised "
          "experimental C_L and C_M loops for held-out NACA 0012 frames [NASA TM-84245]. "
          "Bottom row: AMES-01 cross-check.")
P(f"Across the {_vs['NACA0012 held-out frames']} held-out NACA 0012 frames the frozen "
  f"model predicts peak lift to a mean error of {_vs['mean |CLmax| error [%]']} % and the "
  f"pitching-moment break to a mean error of {_vs['mean |CMmin| error [abs]']} — strong "
  "agreement with real experiment for a reduced-order model, with no tuning to the "
  "held-out data. The mean RMS over the whole C_L loop is "
  f"{_vs['mean RMS_CL']}: the integral peaks are matched considerably better than the "
  "full loop shape, the residual being dominated by the downstroke / reattachment "
  "branch.")
P("Resolution of the damping metric. Recomputing the normalised damping "
  "Ξ̂ from the measured C_M loops and from the model at the same conditions "
  f"gives a mean discrepancy of {_vs['mean |Xi_hat| model-exp discrepancy']} "
  f"(max {_vs['max |Xi_hat| model-exp discrepancy']}). The solver therefore reports "
  f"“neutral” for any |Ξ̂| below "
  f"{_vs['solver damping neutral band (us.DAMPING_TOL)']}, and both cases of this study "
  "fall well inside it. Per-frame Ξ̂ values are in validation_nasa_real.csv.")

H("12.3 Path to formal certification", 2)
P("A drop-in harness (validate_digitized.py) is provided: place a CSV of digitised "
  "experimental points (α, C_L, C_M, C_D) from a named figure of TP-1100/TM-84245 "
  "into 06_postprocessing/validation/experimental/ with its condition row, and the "
  "harness runs the frozen model at that condition and writes RMS C_L, RMS C_M, "
  "max |ΔC_L|, peak-lift error, moment-break error and lift-loop-area error, with an "
  "experiment-vs-model overlay figure — ready for the specific loops of record "
  "required for certification. With no data present it writes a template and exits "
  "without fabricating anything.")

# ================================================================ 13 FIELDS
H("13. Pressure, Temperature, Vector and 3-D Fields", 1)
plots = ROOT/"06_postprocessing"/"plots"
groups = [
 ("Pressure-coefficient contours", "contour_Cp_"),
 ("Velocity magnitude + streamlines", "contour_speed_stream_"),
 ("Velocity vector fields", "contour_vectors_"),
 ("Vorticity (dynamic-stall vortex)", "contour_vorticity_"),
 ("Local Mach-number contours", "contour_Mach_"),
 ("Static air-temperature contours", "contour_Tstatic_"),
 ("Recovery (skin) temperature contours", "contour_Trecovery_"),
]
for _j, (title, pref) in enumerate(groups, start=1):
    H(f"13.{_j} {title}", 2)
    if pref == "contour_vorticity_":
        P("The dynamic-stall vortex is a roll-up of upper-surface boundary-layer "
          "vorticity and so rotates in the same sense as the bound circulation; both "
          "appear with the same sign below. The suction it produces comes from its "
          "low-pressure core, and the reversed flow beneath the core as it convects "
          "aft is the characteristic signature of the stall. The reported loads come "
          "from the UIBS core and do not depend on this reconstruction.",
          italic=True, size=10)
    for f in sorted(plots.glob(pref + "*.png")):
        add_image(f, 5.4, f.stem.replace("_", " "))
H(f"13.{len(groups)+1} Surface recovery-temperature profiles", 2)
for f in sorted(plots.glob("temperature_profile_*.png")):
    add_image(f, 5.8, f.stem.replace("_", " "))
H(f"13.{len(groups)+2} Three-dimensional outputs", 2)
for f in sorted(plots.glob("fig3d_*.png")):
    add_image(f, 5.8, f.stem.replace("_", " "))
P("Response-surface data (head):", bold=True)
_rsdf = pd.read_csv(ROOT/"05_solution"/"response_surface.csv")
_n_ex = int((~_rsdf["within_calibration"]).sum())
_a_cal = float(pd.read_csv(ROOT/'03_model_setup'/'static_polar_reference.csv')['alpha_deg'].max())
add_table_from_df(_rsdf, max_rows=12)
P(f"The grid is bounded by the calibration rather than by preference. The separation law "
  f"f(α) is fitted by inverse Kirchhoff to a static polar tabulated to {_a_cal:.0f}°, so with "
  f"the amplitude held at the case-A value the mean incidence stops at "
  f"{_rsdf['alpha_mean_deg'].max():.1f}° and the peak never exceeds {_rsdf['peak_alpha_deg'].max():.0f}°. "
  f"All {len(_rsdf)} points therefore rest on calibrated data ({_n_ex} extrapolated), which the "
  "within_calibration column records and the build asserts. An earlier version of this sweep "
  "ran to a 26° peak, placing half its points on the extrapolated branch of f(α); those points "
  "are no longer computed, because an uncalibrated result is not made safe by a label. Both "
  "reported cases (§11) peak at exactly the top of the calibration range.",
  italic=True, size=10)

# ================================================================ 14 DRAWINGS
H("14. Engineering Drawings", 1)
P("Dimensioned engineering drawings of the case-study aircraft (third-angle "
  "projection, mm, ISO-style title block).")
dr = ROOT/"08_engineering_drawings"
for f, capt in [("sheet1_general_arrangement_3view.png", "Sheet 1 — general-arrangement three-view."),
                ("sheet2_isometric.png", "Sheet 2 — isometric general arrangement."),
                ("sheet3_main_rotor_blade.png", "Sheet 3 — main-rotor blade."),
                ("sheet4_section_AA_airfoil.png", "Sheet 4 — section A-A (NACA 0012).")]:
    add_image(dr/f, 6.6, capt)

# ================================================================ 15 INVENTORY
H("15. Data Inventory (all generated files)", 1)
# The folder list used to be hardcoded, and it omitted
# 06_postprocessing/validation/experimental/, so two files the pipeline writes
# (conditions.csv and TEMPLATE_experiment.csv) were missing from a section whose
# title promises ALL generated files. Scanned instead, so a new output folder
# cannot silently fall outside the inventory. 07_report is excluded: its
# intermediates are rebuilt every run and are not part of the shipped data.
inv = []
_subs = sorted({str(f.parent.relative_to(ROOT))
                for pat in ("*.csv", "*.png", "*.json")
                for f in ROOT.glob("0[1-8]_*/**/" + pat)
                if "07_report" not in str(f)})
for sub in _subs:
    d = ROOT/sub
    for f in sorted(d.glob("*.csv")):
        try: rows = sum(1 for _ in open(f))-1
        except: rows = "?"
        inv.append([sub, f.name, rows])
    for f in sorted(d.glob("*.png")):
        inv.append([sub, f.name, "-"])
    for f in sorted(d.glob("*.json")):
        inv.append([sub, f.name, "-"])
# The "type" column was dropped: it restated what the extension already says
# (csv / png / json) while consuming width this table cannot spare. With a
# 41-character folder path and a 45-character filename both present, the frame
# could not hold four columns without breaking one of them mid-word.
add_table_from_df(pd.DataFrame(inv, columns=["folder","file","rows"]),
                  max_rows=200)

# ================================================================ 16 SOURCES
H("16. Data Provenance, Validation & Calibration Sources", 1)
P("Honesty note: the static comparison (§12.1) is a calibration check (the static "
  "separation law is fit to the published polar). For the dynamic case the constants "
  "are calibrated on ONE real NACA 0012 loop (frame 9302) and then FROZEN; the frozen "
  f"model PREDICTS {_vs['NACA0012 held-out frames']} held-out real NACA 0012 loops "
  f"(NASA TM-84245) with mean peak-lift error {_vs['mean |CLmax| error [%]']} % and mean "
  f"moment-break error {_vs['mean |CMmin| error [abs]']} (§12.2) — genuine prediction, not "
  "fitting. Airfoil identity is CONFIRMED from the source repository's load_frame.m "
  "mapping (frames 7019–14220 = NACA 0012), so there is no airfoil ambiguity. A drop-in "
  "harness (§12.3) is provided for the specific digitised loops required for formal "
  "certification. Digitised-data source: L. Pancini, BL-DSM-JFS-2021 repository "
  "(original data NASA TM-84245).", italic=True)
P("Licensing of that data. The six frame_*.mat files this study redistributes in "
  "06_postprocessing/validation/experimental/nasa_frames/, and the exp_frame_*.csv "
  "extracts taken from them, are third-party data and are NOT covered by the CC BY "
  "4.0 grant that covers the rest of this repository — they are not the author's to "
  "license. The original measurements are NASA TM-84245, a work of the U.S. "
  "Government; the files were obtained from the BL-DSM-JFS-2021 repository, which "
  "states no licence of its own. See NOTICE and that directory's PROVENANCE.txt.",
  italic=True, size=10)
for s in [
 "[S1] Sheldahl, R.E. & Klimas, P.C. (1981). Aerodynamic Characteristics of Seven "
 "Symmetrical Airfoil Sections..., SAND80-2114, Sandia National Laboratories.",
 "[S2] Abbott, I.H. & von Doenhoff, A.E. (1959). Theory of Wing Sections. Dover.",
 "[S3] McCroskey, W.J. (1987). A Critical Assessment of Wind-Tunnel Results for the "
 "NACA 0012 Airfoil, NASA TM-100019.",
 "[S4] McAlister, K.W., Carr, L.W. & McCroskey, W.J. (1978). Dynamic Stall "
 "Experiments on the NACA 0012 Airfoil, NASA TP-1100.",
 "[S5] McCroskey, W.J. et al. (1982). An Experimental Study of Dynamic Stall on "
 "Advanced Airfoil Sections, NASA TM-84245.",
 "[S6] Leishman, J.G. (2006). Principles of Helicopter Aerodynamics, 2nd ed., "
 "Cambridge University Press.",
 "[S7] Leishman, J.G. & Beddoes, T.S. (1989). A Semi-Empirical Model for Dynamic "
 "Stall, J. American Helicopter Society, 34(3).",
 "[S8] Carr, L.W. (1988). Progress in Analysis and Prediction of Dynamic Stall, "
 "J. Aircraft, 25(1)."]:
    p = doc.add_paragraph(); r = p.add_run(s); r.font.size = Pt(10); r.font.color.rgb = INK

# ---- document metadata (author only; no third-party attribution) ----
cp = doc.core_properties
cp.author = AUTHOR; cp.last_modified_by = AUTHOR; cp.title = \
    "Dynamic-Stall Prediction Case Study — UNISTALL Universal Solver"
cp.comments = "Authored by " + AUTHOR

OUT = ROOT/"07_report"/"case.docx"
doc.save(str(OUT))
# also place a copy at project root for convenience
import shutil; shutil.copy(str(OUT), str(ROOT/"case.docx"))
print("[docx] wrote", OUT, "and", ROOT/"case.docx")
